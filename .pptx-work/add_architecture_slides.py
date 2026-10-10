from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

SOURCE = "DermTrack_Project_Review_With_Proposed_Solution.pptx"
OUT = "DermTrack_Project_Review_With_Architecture.pptx"
prs = Presentation(SOURCE)

navy = RGBColor(10, 25, 47)
navy2 = RGBColor(16, 42, 67)
teal = RGBColor(35, 194, 176)
coral = RGBColor(255, 126, 103)
white = RGBColor(246, 250, 252)
muted = RGBColor(181, 204, 213)
soft = RGBColor(222, 237, 240)
line = RGBColor(52, 91, 111)
rose = RGBColor(54, 49, 64)
rose_line = RGBColor(111, 82, 91)


def textbox(slide, text, x, y, w, h, size=16, color=white, bold=False,
            font="Aptos", align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=0):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.clear(); tf.margin = margin; tf.vertical_anchor = valign; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = text; p.alignment = align
    p.font.name = font; p.font.size = Pt(size); p.font.color.rgb = color; p.font.bold = bold
    return box


def header(slide, title, number):
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = navy
    textbox(slide, "PRELIMINARY ARCHITECTURE", 0.72, 0.46, 5.0, 0.25, 10, teal, True)
    textbox(slide, title, 0.72, 0.78, 9.8, 0.6, 28, white, True, "Aptos Display")
    textbox(slide, f"0{number}  /  DERMTRACK AGENT", 10.1, 0.52, 2.5, 0.25, 9, muted, True, align=PP_ALIGN.RIGHT)
    a = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.72), Inches(1.53), Inches(0.72), Inches(0.06))
    a.fill.solid(); a.fill.fore_color.rgb = coral; a.line.fill.background()


def data_card(slide, x, y, w, h, title, subtitle, body, accent):
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    box.fill.solid(); box.fill.fore_color.rgb = navy2; box.line.color.rgb = line; box.line.width = Pt(1)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(0.09), Inches(h))
    bar.fill.solid(); bar.fill.fore_color.rgb = accent; bar.line.fill.background()
    textbox(slide, title, x+0.30, y+0.25, w-0.55, 0.28, 16, accent, True)
    textbox(slide, subtitle, x+0.30, y+0.65, w-0.55, 0.25, 11, soft, True)
    textbox(slide, body, x+0.30, y+1.08, w-0.55, h-1.25, 12, white, False)

# Slide 11: image placeholder only
slide = prs.slides.add_slide(prs.slide_layouts[6])
header(slide, "System architecture diagram", 11)
# Intentionally blank below the title so the user can paste the generated architecture image.

# Slide 12: database/data-management explanation
slide = prs.slides.add_slide(prs.slide_layouts[6])
header(slide, "Role of the database and data-management layer", 12)
textbox(slide, "DermTrack separates operational patient data from searchable medical knowledge so each data type can be stored and accessed appropriately.", 0.72, 1.69, 11.7, 0.38, 13, muted)

data_card(slide, 0.72, 2.30, 3.72, 2.40, "MONGODB ATLAS", "Operational application database", "Stores structured patient intake, consultation records, transcripts, image observations, triage outcomes, retrieved-source metadata, generated responses, timestamps, and history.", teal)
data_card(slide, 4.80, 2.30, 3.72, 2.40, "GRIDFS", "Large-file image storage", "Stores uploaded JPEG, PNG, and WEBP consultation images separately from the consultation document. The record keeps the GridFS file ID and content type.", coral)
data_card(slide, 8.88, 2.30, 3.72, 2.40, "CHROMADB OR ATLAS VECTOR SEARCH", "Medical knowledge retrieval", "Stores trusted-document chunks, embeddings, source URLs, trust tiers, page numbers, and document metadata. It is not used as the patient-record store.", teal)

# Data lifecycle strip
strip = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.72), Inches(5.12), Inches(11.90), Inches(0.92))
strip.fill.solid(); strip.fill.fore_color.rgb = RGBColor(20, 67, 80); strip.line.fill.background()
textbox(slide, "DATA LIFECYCLE", 1.02, 5.39, 1.55, 0.22, 10, coral, True)
textbox(slide, "Collect → Validate → Process → Retrieve evidence → Generate response → Persist consultation", 2.75, 5.35, 9.20, 0.28, 14, white, True)
textbox(slide, "Design principle: patient data remains in the application database, while medical knowledge is indexed separately for retrieval.", 0.72, 6.40, 11.85, 0.25, 10, muted, False, align=PP_ALIGN.CENTER)

prs.core_properties.title = "DermTrack Agent — Preliminary Architecture"
prs.core_properties.subject = "Architecture placeholder and data-management role"
prs.save(OUT)
print(OUT)
