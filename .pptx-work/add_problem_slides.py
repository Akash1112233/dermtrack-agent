from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

SOURCE = "DermTrack_Introduction_Slide.pptx"
OUT = "DermTrack_Project_Review_Intro_Problem.pptx"
prs = Presentation(SOURCE)

# Shared palette
navy = RGBColor(10, 25, 47)
navy2 = RGBColor(16, 42, 67)
teal = RGBColor(35, 194, 176)
coral = RGBColor(255, 126, 103)
white = RGBColor(246, 250, 252)
muted = RGBColor(181, 204, 213)
line = RGBColor(52, 91, 111)
soft = RGBColor(222, 237, 240)


def textbox(slide, text, x, y, w, h, size=16, color=white, bold=False,
            font="Aptos", align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=0):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.clear(); tf.margin = margin; tf.vertical_anchor = valign; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = text; p.alignment = align
    p.font.name = font; p.font.size = Pt(size); p.font.color.rgb = color; p.font.bold = bold
    return box


def rounded_card(slide, x, y, w, h, fill=navy2, border=line):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    card.fill.solid(); card.fill.fore_color.rgb = fill
    card.line.color.rgb = border; card.line.width = Pt(1)
    return card


def header(slide, section, title, number):
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = navy
    textbox(slide, section.upper(), 0.72, 0.46, 3.8, 0.25, 10, teal, True)
    textbox(slide, title, 0.72, 0.78, 9.3, 0.6, 30, white, True, "Aptos Display")
    textbox(slide, f"0{number}  /  DERMTRACK AGENT", 10.1, 0.52, 2.5, 0.25, 9, muted, True, align=PP_ALIGN.RIGHT)
    # small corner accent
    a = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.72), Inches(1.53), Inches(0.72), Inches(0.06))
    a.fill.solid(); a.fill.fore_color.rgb = coral; a.line.fill.background()

# Slide 2: Problem definition
slide = prs.slides.add_slide(prs.slide_layouts[6])
header(slide, "Problem identification", "The gap DermTrack addresses", 2)

# Main statement
rounded_card(slide, 0.72, 1.92, 6.15, 2.02)
textbox(slide, "THE CORE PROBLEM", 1.05, 2.20, 2.2, 0.25, 10, coral, True)
textbox(slide, "People struggle to interpret skin symptoms and decide what to do next. Online information is often fragmented, hard to verify, and unsafe without context.", 1.05, 2.58, 5.45, 1.10, 16, white, True)

# Four pain points on right
textbox(slide, "Why this is difficult", 7.35, 1.93, 4.5, 0.35, 18, white, True)
pain_points = [
    ("01", "Unstructured inputs", "Symptoms, history, and images arrive in different formats."),
    ("02", "Evidence overload", "Reliable medical information is scattered across sources."),
    ("03", "Safety risk", "A confident but unsupported answer can delay appropriate care."),
    ("04", "Limited continuity", "Previous consultations and observations are hard to organize."),
]
for i, (num, title, desc) in enumerate(pain_points):
    y = 2.42 + i * 0.72
    textbox(slide, num, 7.36, y, 0.40, 0.28, 11, teal, True)
    textbox(slide, title, 7.90, y-0.02, 2.35, 0.25, 14, white, True)
    textbox(slide, desc, 7.90, y+0.25, 4.45, 0.30, 10.5, muted)

# Intended response strip
strip = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.72), Inches(4.42), Inches(11.9), Inches(1.35))
strip.fill.solid(); strip.fill.fore_color.rgb = RGBColor(20, 67, 80); strip.line.fill.background()
textbox(slide, "DERMTRACK'S RESPONSE", 1.05, 4.72, 2.2, 0.22, 10, teal, True)
textbox(slide, "Structured intake + optional image observations + safety triage + trusted retrieval + consultation history.", 3.20, 4.65, 8.85, 0.54, 14, white, True)
textbox(slide, "Supports informed next steps; it does not replace a dermatologist or issue prescriptions.", 3.20, 5.25, 8.85, 0.25, 10, soft, False)

# Slide 3: user/domain/motivation/relevance/impact
slide = prs.slides.add_slide(prs.slide_layouts[6])
header(slide, "Problem identification", "Who needs it—and why it matters", 3)

cards = [
    (0.72, 1.88, 3.65, 2.04, "TARGET USERS", "Patients seeking understandable guidance\n\nClinicians reviewing structured consultation context", teal),
    (4.84, 1.88, 3.65, 2.04, "DOMAIN", "AI-assisted dermatology support\n\nText + optional skin images + medical knowledge retrieval", coral),
    (8.96, 1.88, 3.65, 2.04, "MOTIVATION", "Make first-line information more organized, explainable, privacy-conscious, and safety-aware", teal),
]
for x, y, w, h, label, body, accent in cards:
    rounded_card(slide, x, y, w, h)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(0.09), Inches(h))
    bar.fill.solid(); bar.fill.fore_color.rgb = accent; bar.line.fill.background()
    textbox(slide, label, x+0.28, y+0.28, w-0.50, 0.25, 10, accent, True)
    textbox(slide, body, x+0.28, y+0.68, w-0.50, h-0.88, 14, white, True)

# Practical relevance and impact lower section
rounded_card(slide, 0.72, 4.35, 5.72, 1.55, fill=RGBColor(20, 67, 80))
textbox(slide, "PRACTICAL RELEVANCE", 1.05, 4.65, 2.5, 0.23, 10, teal, True)
textbox(slide, "A realistic workflow for collecting better consultation context, grounding answers in trusted sources, flagging warning signs, and preserving a usable history.", 1.05, 5.02, 4.95, 0.55, 15, white, True)

rounded_card(slide, 6.76, 4.35, 5.85, 1.55, fill=RGBColor(54, 49, 64), border=RGBColor(111, 82, 91))
textbox(slide, "EXPECTED IMPACT", 7.09, 4.65, 2.5, 0.23, 10, coral, True)
textbox(slide, "Improve access to understandable, evidence-linked support; encourage earlier professional attention when red flags appear; and reduce unsupported, overconfident health responses.", 7.09, 5.02, 5.05, 0.55, 15, white, True)

textbox(slide, "Expected impact is a design objective, not a claim of clinical effectiveness; validation with clinicians and real-world users remains future work.", 0.74, 6.45, 11.7, 0.28, 10, muted, False, align=PP_ALIGN.CENTER)

prs.core_properties.title = "DermTrack Agent — Introduction and Problem Identification"
prs.core_properties.subject = "Project review slides"
prs.save(OUT)
print(OUT)
