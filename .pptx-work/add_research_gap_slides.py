from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

SOURCE = "DermTrack_Project_Review_With_Literature.pptx"
OUT = "DermTrack_Project_Review_With_Research_Gap.pptx"
prs = Presentation(SOURCE)

navy = RGBColor(10, 25, 47)
navy2 = RGBColor(16, 42, 67)
teal = RGBColor(35, 194, 176)
coral = RGBColor(255, 126, 103)
white = RGBColor(246, 250, 252)
muted = RGBColor(181, 204, 213)
line = RGBColor(52, 91, 111)
soft = RGBColor(222, 237, 240)
rose = RGBColor(54, 49, 64)
rose_line = RGBColor(111, 82, 91)


def textbox(slide, text, x, y, w, h, size=16, color=white, bold=False,
            font="Aptos", align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=0):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.clear(); tf.margin = margin; tf.vertical_anchor = valign; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = text; p.alignment = align
    p.font.name = font; p.font.size = Pt(size); p.font.color.rgb = color; p.font.bold = bold
    return box


def header(slide, title, number):
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = navy
    textbox(slide, "RESEARCH GAP", 0.72, 0.46, 3.5, 0.25, 10, teal, True)
    textbox(slide, title, 0.72, 0.78, 9.8, 0.6, 28, white, True, "Aptos Display")
    textbox(slide, f"0{number}  /  DERMTRACK AGENT", 10.1, 0.52, 2.5, 0.25, 9, muted, True, align=PP_ALIGN.RIGHT)
    a = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.72), Inches(1.53), Inches(0.72), Inches(0.06))
    a.fill.solid(); a.fill.fore_color.rgb = coral; a.line.fill.background()


def gap_card(slide, x, y, w, h, num, title, body, accent=teal):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    card.fill.solid(); card.fill.fore_color.rgb = navy2; card.line.color.rgb = line; card.line.width = Pt(1)
    circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x+0.25), Inches(y+0.22), Inches(0.46), Inches(0.46))
    circle.fill.solid(); circle.fill.fore_color.rgb = accent; circle.line.fill.background()
    textbox(slide, num, x+0.25, y+0.295, 0.46, 0.18, 9, navy, True, align=PP_ALIGN.CENTER)
    textbox(slide, title, x+0.88, y+0.25, w-1.13, 0.30, 14, white, True)
    textbox(slide, body, x+0.28, y+0.82, w-0.56, h-1.02, 11.5, muted)

# Slide 7: research gaps
slide = prs.slides.add_slide(prs.slide_layouts[6])
header(slide, "What existing approaches still leave unresolved", 7)
textbox(slide, "The literature shows strong progress, but many systems solve only one part of the dermatology-support problem.", 0.72, 1.69, 11.7, 0.34, 13, muted)

gap_card(slide, 0.72, 2.25, 3.72, 1.55, "01", "Image-only focus", "Many systems classify images but do not combine symptoms, history, medications, and patient context in one workflow.", teal)
gap_card(slide, 4.80, 2.25, 3.72, 1.55, "02", "Unsafe generation", "Language models can produce fluent but unsupported medical content, even when the answer sounds confident.", coral)
gap_card(slide, 8.88, 2.25, 3.72, 1.55, "03", "Benchmark gap", "High accuracy on static datasets does not prove safe triage, calibration, usability, or clinical effectiveness.", teal)
gap_card(slide, 0.72, 4.15, 3.72, 1.55, "04", "Representation gap", "Skin-tone imbalance, label quality, and dataset shift can reduce reliability for underrepresented groups.", coral)
gap_card(slide, 4.80, 4.15, 3.72, 1.55, "05", "Evidence gap", "Retrieval can fail through irrelevant, incomplete, stale, or weakly sourced medical information.", teal)
gap_card(slide, 8.88, 4.15, 3.72, 1.55, "06", "Continuity gap", "Many prototypes do not preserve structured, longitudinal consultation history for follow-up context.", coral)
textbox(slide, "Research gap: there is a need for a bounded end-to-end support workflow—not just a better image classifier or a larger language model.", 0.72, 6.36, 11.85, 0.34, 12.5, white, True, align=PP_ALIGN.CENTER)

# Slide 8: gap-to-design mapping
slide = prs.slides.add_slide(prs.slide_layouts[6])
header(slide, "DermTrack's response to the research gap", 8)
textbox(slide, "DermTrack addresses the gap by integrating complementary components while keeping the system non-diagnostic and human-supervised.", 0.72, 1.69, 11.7, 0.34, 13, muted)

# Mapping table-like rows
rows = [
    ("Patient context", "Structured intake + symptom/history fields", "Moves beyond image-only analysis", teal),
    ("Grounded answers", "Local trusted RAG + source metadata", "Reduces unsupported medical generation", coral),
    ("Current information", "Conditional Tavily research with de-identified queries", "Adds bounded, optional web evidence", teal),
    ("Safety first", "Red-flag triage + explicit non-diagnostic boundaries", "Prioritizes escalation over false reassurance", coral),
    ("Continuity", "MongoDB consultation records + GridFS images", "Preserves organized longitudinal context", teal),
    ("Responsible deployment", "Human review, privacy controls, and audit-ready state", "Supports evaluation before clinical claims", coral),
]
start_y = 2.23
for i, (gap, design, contribution, accent) in enumerate(rows):
    y = start_y + i * 0.62
    fill = navy2 if i % 2 == 0 else RGBColor(19, 51, 76)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.72), Inches(y), Inches(0.09), Inches(0.52))
    bar.fill.solid(); bar.fill.fore_color.rgb = accent; bar.line.fill.background()
    back = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.81), Inches(y), Inches(11.80), Inches(0.52))
    back.fill.solid(); back.fill.fore_color.rgb = fill; back.line.color.rgb = line; back.line.width = Pt(0.5)
    textbox(slide, gap, 1.05, y+0.13, 2.15, 0.22, 11.5, accent, True)
    textbox(slide, design, 3.45, y+0.11, 3.85, 0.27, 11.5, white, True)
    textbox(slide, contribution, 7.75, y+0.11, 4.45, 0.27, 11.5, muted)

# Bottom boundary statement
boundary = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.72), Inches(6.18), Inches(11.90), Inches(0.58))
boundary.fill.solid(); boundary.fill.fore_color.rgb = rose; boundary.line.color.rgb = rose_line
textbox(slide, "Contribution boundary: DermTrack is a research prototype for safer information support—not a validated diagnostic or treatment system.", 1.02, 6.34, 11.30, 0.22, 11.5, white, True, align=PP_ALIGN.CENTER)

prs.core_properties.title = "DermTrack Agent — Research Gap"
prs.core_properties.subject = "Research gap and project contribution slides"
prs.save(OUT)
print(OUT)
