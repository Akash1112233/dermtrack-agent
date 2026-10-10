from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

SOURCE = "DermTrack_Project_Review_With_Research_Gap.pptx"
OUT = "DermTrack_Project_Review_With_Proposed_Solution.pptx"
prs = Presentation(SOURCE)

navy = RGBColor(10, 25, 47)
navy2 = RGBColor(16, 42, 67)
teal = RGBColor(35, 194, 176)
coral = RGBColor(255, 126, 103)
white = RGBColor(246, 250, 252)
muted = RGBColor(181, 204, 213)
line = RGBColor(52, 91, 111)
soft = RGBColor(222, 237, 240)
rose = RGBColor(54, 49, 64)
rose_line = RGBColor(111, 82, 91)


def textbox(slide, text, x, y, w, h, size=16, color=white, bold=False,
            font="Aptos", align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=0):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.clear(); tf.margin = margin; tf.vertical_anchor = valign; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = text; p.alignment = align
    p.font.name = font; p.font.size = Pt(size); p.font.color.rgb = color; p.font.bold = bold
    return box


def header(slide, title, number):
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = navy
    textbox(slide, "PROPOSED SOLUTION", 0.72, 0.46, 4.0, 0.25, 10, teal, True)
    textbox(slide, title, 0.72, 0.78, 9.8, 0.6, 28, white, True, "Aptos Display")
    textbox(slide, f"0{number}  /  DERMTRACK AGENT", 10.1, 0.52, 2.5, 0.25, 9, muted, True, align=PP_ALIGN.RIGHT)
    a = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.72), Inches(1.53), Inches(0.72), Inches(0.06))
    a.fill.solid(); a.fill.fore_color.rgb = coral; a.line.fill.background()


def card(slide, x, y, w, h, title, body, accent=teal):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid(); shp.fill.fore_color.rgb = navy2; shp.line.color.rgb = line; shp.line.width = Pt(1)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(0.08), Inches(h))
    bar.fill.solid(); bar.fill.fore_color.rgb = accent; bar.line.fill.background()
    textbox(slide, title, x+0.28, y+0.22, w-0.48, 0.25, 11, accent, True)
    textbox(slide, body, x+0.28, y+0.62, w-0.48, h-0.82, 13, white, True)

# Slide 9: solution overview
slide = prs.slides.add_slide(prs.slide_layouts[6])
header(slide, "A bounded, grounded, human-supervised support system", 9)
textbox(slide, "DermTrack addresses the identified gap by combining multimodal intake, safety triage, trusted evidence retrieval, and persistent consultation context in one workflow.", 0.72, 1.69, 11.7, 0.42, 13, muted)

card(slide, 0.72, 2.35, 3.72, 1.65, "INPUT LAYER", "Structured patient details\nOptional image upload\nAudio transcription through Deepgram", teal)
card(slide, 4.80, 2.35, 3.72, 1.65, "REASONING LAYER", "Gemini image observations\nSafety red-flag triage\nLangGraph orchestration", coral)
card(slide, 8.88, 2.35, 3.72, 1.65, "EVIDENCE LAYER", "Local RAG from trusted sources\nChromaDB or Atlas vector search\nConditional Tavily research", teal)

# arrows between layers
for x in (4.48, 8.56):
    arr = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(3.15), Inches(x+0.28), Inches(3.15))
    arr.line.color.rgb = coral; arr.line.width = Pt(2.2); arr.line.end_arrowhead = True

card(slide, 0.72, 4.50, 3.72, 1.55, "SAFETY BOUNDARY", "Non-diagnostic language\nNo invented prescriptions\nEscalation when warning signs appear", coral)
card(slide, 4.80, 4.50, 3.72, 1.55, "RESPONSE LAYER", "Evidence-linked educational response\nUncertainty-aware wording\nHuman professional remains responsible", teal)
card(slide, 8.88, 4.50, 3.72, 1.55, "PERSISTENCE LAYER", "MongoDB consultation records\nGridFS image storage\nLongitudinal patient context", coral)
textbox(slide, "The proposed solution is an integrated research prototype, not a replacement for clinical diagnosis or treatment.", 0.72, 6.40, 11.85, 0.25, 10, muted, False, align=PP_ALIGN.CENTER)

# Slide 10: workflow
slide = prs.slides.add_slide(prs.slide_layouts[6])
header(slide, "End-to-end solution workflow", 10)
textbox(slide, "Each stage is bounded, testable, and designed to preserve safety and evidence provenance.", 0.72, 1.69, 11.7, 0.34, 13, muted)

steps = [
    ("1", "Collect", "Patient intake\nText / audio / image", teal),
    ("2", "Normalize", "Validate fields\nTranscribe audio", coral),
    ("3", "Observe", "Analyze image\nDescribe visible features", teal),
    ("4", "Triage", "Check red flags\nEscalate urgent cases", coral),
    ("5", "Ground", "Retrieve trusted sources\nResearch only when needed", teal),
    ("6", "Respond", "Generate bounded answer\nPersist consultation", coral),
]
start_x = 0.72
for i,(num,title,body,accent) in enumerate(steps):
    x = start_x + i*2.04
    c = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(2.45), Inches(1.72), Inches(1.90))
    c.fill.solid(); c.fill.fore_color.rgb = navy2; c.line.color.rgb = line; c.line.width = Pt(1)
    circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x+0.60), Inches(2.67), Inches(0.52), Inches(0.52))
    circle.fill.solid(); circle.fill.fore_color.rgb = accent; circle.line.fill.background()
    textbox(slide, num, x+0.60, 2.80, 0.52, 0.18, 11, navy, True, align=PP_ALIGN.CENTER)
    textbox(slide, title, x+0.15, 3.38, 1.42, 0.25, 14, white, True, align=PP_ALIGN.CENTER)
    textbox(slide, body, x+0.15, 3.76, 1.42, 0.42, 10.5, muted, False, align=PP_ALIGN.CENTER)
    if i < len(steps)-1:
        arr = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x+1.72), Inches(3.40), Inches(x+1.98), Inches(3.40))
        arr.line.color.rgb = teal; arr.line.width = Pt(1.6); arr.line.end_arrowhead = True

# architecture summary band
band = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.72), Inches(4.85), Inches(11.90), Inches(1.10))
band.fill.solid(); band.fill.fore_color.rgb = RGBColor(20, 67, 80); band.line.fill.background()
textbox(slide, "IMPLEMENTATION STACK", 1.02, 5.10, 2.1, 0.23, 10, teal, True)
textbox(slide, "FastAPI  ·  LangGraph  ·  Gemini Vision  ·  Deepgram  ·  ChromaDB / Atlas  ·  Tavily  ·  MongoDB + GridFS", 3.05, 5.07, 8.9, 0.25, 13, white, True)
textbox(slide, "Research contribution: a single workflow connects multimodal context, safety checks, evidence retrieval, and longitudinal persistence.", 0.72, 6.40, 11.85, 0.25, 10, muted, False, align=PP_ALIGN.CENTER)

prs.core_properties.title = "DermTrack Agent — Proposed Solution"
prs.core_properties.subject = "Proposed solution and workflow slides"
prs.save(OUT)
print(OUT)
