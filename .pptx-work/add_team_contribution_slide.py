from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

SOURCE = "DermTrack_Project_Review_With_Architecture.pptx"
OUT = "DermTrack_Project_Review_With_Team_Contribution.pptx"
prs = Presentation(SOURCE)

navy = RGBColor(10, 25, 47)
navy2 = RGBColor(16, 42, 67)
teal = RGBColor(35, 194, 176)
coral = RGBColor(255, 126, 103)
white = RGBColor(246, 250, 252)
muted = RGBColor(181, 204, 213)
line = RGBColor(52, 91, 111)
soft = RGBColor(222, 237, 240)
rose = RGBColor(54, 49, 64)
rose_line = RGBColor(111, 82, 91)


def textbox(slide, text, x, y, w, h, size=16, color=white, bold=False,
            font="Aptos", align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=0):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.clear(); tf.margin = margin; tf.vertical_anchor = valign; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = text; p.alignment = align
    p.font.name = font; p.font.size = Pt(size); p.font.color.rgb = color; p.font.bold = bold
    return box


def header(slide, title, number):
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = navy
    textbox(slide, "TEAM CONTRIBUTION", 0.72, 0.46, 4.0, 0.25, 10, teal, True)
    textbox(slide, title, 0.72, 0.78, 9.8, 0.6, 28, white, True, "Aptos Display")
    textbox(slide, f"0{number}  /  DERMTRACK AGENT", 10.1, 0.52, 2.5, 0.25, 9, muted, True, align=PP_ALIGN.RIGHT)
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.72), Inches(1.53), Inches(0.72), Inches(0.06))
    accent.fill.solid(); accent.fill.fore_color.rgb = coral; accent.line.fill.background()


def member_card(slide, x, name, roll, accent, role, items):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(2.08), Inches(5.72), Inches(3.75))
    card.fill.solid(); card.fill.fore_color.rgb = navy2; card.line.color.rgb = line; card.line.width = Pt(1)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(2.08), Inches(0.10), Inches(3.75))
    bar.fill.solid(); bar.fill.fore_color.rgb = accent; bar.line.fill.background()
    textbox(slide, name, x+0.35, 2.38, 4.9, 0.35, 22, white, True, "Aptos Display")
    textbox(slide, roll, x+0.35, 2.82, 2.2, 0.25, 11, accent, True)
    textbox(slide, role, x+0.35, 3.25, 4.9, 0.28, 13, soft, True)
    y = 3.73
    for item in items:
        dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x+0.38), Inches(y+0.06), Inches(0.10), Inches(0.10))
        dot.fill.solid(); dot.fill.fore_color.rgb = accent; dot.line.fill.background()
        textbox(slide, item, x+0.62, y, 4.72, 0.30, 11.5, muted)
        y += 0.43

slide = prs.slides.add_slide(prs.slide_layouts[6])
header(slide, "Equal technical responsibilities", 13)
textbox(slide, "Both members contribute to the core implementation, not only documentation or presentation work.", 0.72, 1.69, 11.7, 0.34, 13, muted)

member_card(slide, 0.72, "Annam Akash", "23MIA1112", teal, "Backend, AI workflow & data engineering", [
    "Design FastAPI APIs and LangGraph workflow orchestration.",
    "Implement multimodal input handling and Gemini/Deepgram integration.",
    "Develop RAG ingestion, chunking, embeddings, and vector retrieval.",
    "Implement MongoDB persistence, GridFS image storage, and consultation history.",
])
member_card(slide, 6.76, "B. Hema", "23MIA1101", coral, "Client experience, evaluation & responsible AI", [
    "Develop the patient/clinician interface and structured intake experience.",
    "Implement response presentation, source citations, and history views.",
    "Design evaluation datasets, safety tests, usability checks, and error analysis.",
    "Support privacy, fairness, documentation, and end-to-end integration.",
])

# Shared work footer
footer = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.72), Inches(6.12), Inches(11.90), Inches(0.62))
footer.fill.solid(); footer.fill.fore_color.rgb = rose; footer.line.color.rgb = rose_line
textbox(slide, "SHARED RESPONSIBILITY", 1.02, 6.32, 1.95, 0.20, 10, coral, True)
textbox(slide, "Architecture decisions · Git collaboration · Integration testing · Lecturer review · Final demonstration", 3.15, 6.29, 8.75, 0.22, 11.5, white, True)

prs.core_properties.title = "DermTrack Agent — Team Contribution"
prs.core_properties.subject = "Equal technical responsibility split"
prs.save(OUT)
print(OUT)
