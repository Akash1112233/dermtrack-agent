from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

SOURCE = "DermTrack_Project_Review_With_Team_Contribution.pptx"
OUT = "DermTrack_Project_Review_Complete_Review1.pptx"
prs = Presentation(SOURCE)

navy = RGBColor(10, 25, 47)
navy2 = RGBColor(16, 42, 67)
teal = RGBColor(35, 194, 176)
coral = RGBColor(255, 126, 103)
white = RGBColor(246, 250, 252)
muted = RGBColor(181, 204, 213)
line = RGBColor(52, 91, 111)
soft = RGBColor(222, 237, 240)
rose = RGBColor(54, 49, 64)
rose_line = RGBColor(111, 82, 91)

def textbox(slide, text, x, y, w, h, size=16, color=white, bold=False,
            font="Aptos", align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=0):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.clear(); tf.margin = margin; tf.vertical_anchor = valign; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = text; p.alignment = align
    p.font.name = font; p.font.size = Pt(size); p.font.color.rgb = color; p.font.bold = bold
    return box

def header(slide, section, title, number):
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = navy
    textbox(slide, section.upper(), 0.72, 0.46, 5.0, 0.25, 10, teal, True)
    textbox(slide, title, 0.72, 0.78, 10.0, 0.6, 28, white, True, "Aptos Display")
    textbox(slide, f"{number:03d}  /  DERMTRACK AGENT", 10.1, 0.52, 2.5, 0.25, 9, muted, True, align=PP_ALIGN.RIGHT)
    # No decorative underline: use a small section marker instead.
    marker = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.72), Inches(1.53), Inches(0.72), Inches(0.06))
    marker.fill.solid(); marker.fill.fore_color.rgb = coral; marker.line.fill.background()

def card(slide, x, y, w, h, title, body, accent=teal, body_size=13):
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    box.fill.solid(); box.fill.fore_color.rgb = navy2; box.line.color.rgb = line; box.line.width = Pt(1)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(0.09), Inches(h))
    bar.fill.solid(); bar.fill.fore_color.rgb = accent; bar.line.fill.background()
    textbox(slide, title, x+0.30, y+0.23, w-0.55, 0.28, 12, accent, True)
    textbox(slide, body, x+0.30, y+0.66, w-0.55, h-0.84, body_size, white, False)

def bullet_card(slide, x, y, w, h, title, items, accent=teal):
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    box.fill.solid(); box.fill.fore_color.rgb = navy2; box.line.color.rgb = line; box.line.width = Pt(1)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(0.09), Inches(h))
    bar.fill.solid(); bar.fill.fore_color.rgb = accent; bar.line.fill.background()
    textbox(slide, title, x+0.30, y+0.22, w-0.55, 0.28, 13, accent, True)
    yy=y+0.68
    for item in items:
        dot=slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x+0.34), Inches(yy+0.08), Inches(0.09), Inches(0.09))
        dot.fill.solid(); dot.fill.fore_color.rgb=accent; dot.line.fill.background()
        textbox(slide,item,x+0.58,yy,w-0.85,0.34,11.5,white)
        yy+=0.52

# Append the four missing topic slides; they will be reordered below.
# Objectives
s = prs.slides.add_slide(prs.slide_layouts[6])
header(s, "Review 1 objectives", "Project objectives", 11)
textbox(s, "The objectives translate the research gap into measurable system-development goals.", 0.72, 1.69, 11.7, 0.34, 13, muted)
bullet_card(s, 0.72, 2.25, 5.72, 3.55, "SYSTEM OBJECTIVES", [
    "Build a non-diagnostic dermatology-support prototype.",
    "Accept structured text, audio, and optional skin images.",
    "Provide cautious image observations and red-flag safety triage.",
    "Generate evidence-linked educational responses without prescribing.",
], teal)
bullet_card(s, 6.76, 2.25, 5.72, 3.55, "ENGINEERING OBJECTIVES", [
    "Use LangGraph to coordinate a bounded, testable workflow.",
    "Ground answers with trusted local RAG and optional web research.",
    "Persist consultation context and images for longitudinal follow-up.",
    "Evaluate reliability, safety boundaries, privacy, and usability.",
], coral)
textbox(s, "Success criterion: a working, testable vertical slice with clearly stated limitations—not a claim of clinical validation.", 0.72, 6.35, 11.85, 0.28, 10, muted, False, align=PP_ALIGN.CENTER)

# Novelty
s = prs.slides.add_slide(prs.slide_layouts[6])
header(s, "Expected contribution", "Expected novelty of DermTrack", 12)
textbox(s, "The novelty is in the integration and safety boundary of the workflow, rather than in claiming a new medical diagnosis model.", 0.72, 1.69, 11.7, 0.36, 13, muted)
card(s, 0.72, 2.28, 3.72, 1.62, "01  INTEGRATED CONTEXT", "Combines structured intake, optional images, audio transcription, and consultation history instead of relying on an image alone.", teal, 12)
card(s, 4.80, 2.28, 3.72, 1.62, "02  GROUNDED SUPPORT", "Connects local trusted RAG with source metadata and optional bounded web research for evidence-aware responses.", coral, 12)
card(s, 8.88, 2.28, 3.72, 1.62, "03  SAFETY-FIRST FLOW", "Makes red-flag triage and non-diagnostic boundaries explicit before response generation.", teal, 12)
card(s, 0.72, 4.35, 3.72, 1.62, "04  PRIVACY MINIMIZATION", "Uses de-identified structured queries for optional web research and keeps sensitive consultation data in the application store.", coral, 12)
card(s, 4.80, 4.35, 3.72, 1.62, "05  LONGITUDINAL CONTEXT", "Preserves consultation records and uploaded images to support organized follow-up context.", teal, 12)
card(s, 8.88, 4.35, 3.72, 1.62, "06  EVALUABLE DESIGN", "Separates workflow stages so retrieval, triage, generation, persistence, and limitations can be tested independently.", coral, 12)
textbox(s, "Expected novelty: a privacy-conscious, evidence-grounded support workflow that is explicit about what it can and cannot do.", 0.72, 6.40, 11.85, 0.25, 10, muted, False, align=PP_ALIGN.CENTER)

# Work plan
s = prs.slides.add_slide(prs.slide_layouts[6])
header(s, "Implementation roadmap", "Work plan for the project", 15)
textbox(s, "The work plan moves from evidence and requirements to a tested vertical slice, followed by evaluation and presentation readiness.", 0.72, 1.69, 11.7, 0.36, 13, muted)
phases = [
    ("01", "Foundation", "Requirements, literature survey, research gap, objectives", teal),
    ("02", "Core backend", "FastAPI, schemas, LangGraph workflow, safety triage", coral),
    ("03", "Knowledge + data", "RAG ingestion, vector search, MongoDB, GridFS", teal),
    ("04", "Multimodal + UI", "Audio, image observation, intake interface, history", coral),
    ("05", "Evaluation", "Unit tests, workflow tests, safety checks, error analysis", teal),
    ("06", "Finalization", "Integration demo, documentation, report, lecturer review", coral),
]
for i,(num,title,body,accent) in enumerate(phases):
    x=0.72 + (i%3)*4.08; y=2.35 + (i//3)*1.72
    box=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(3.68), Inches(1.30))
    box.fill.solid(); box.fill.fore_color.rgb=navy2; box.line.color.rgb=line
    circ=s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x+0.22), Inches(y+0.25), Inches(0.52), Inches(0.52))
    circ.fill.solid(); circ.fill.fore_color.rgb=accent; circ.line.fill.background()
    textbox(s,num,x+0.22,y+0.38,0.52,0.18,10,navy,True,align=PP_ALIGN.CENTER)
    textbox(s,title,x+0.93,y+0.22,2.45,0.25,13,white,True)
    textbox(s,body,x+0.93,y+0.58,2.45,0.45,10.5,muted)
textbox(s, "Milestone: demonstrate one complete consultation path before expanding scope.", 0.72, 6.40, 11.85, 0.25, 10, muted, False, align=PP_ALIGN.CENTER)

# Reorder: current original slides 0-12, appended slides 13=obj,14=novelty,15=workplan.
# Desired order: intro/problem/literature/gap/solution/objectives/novelty/architecture/db/workplan/team.
ids = list(prs.slides._sldIdLst)
order = [0,1,2,3,4,5,6,7,8,9,13,14,10,11,15,12]
prs.slides._sldIdLst[:] = [ids[i] for i in order]

# Normalize visible slide numbers for the final ordered deck.
for slide_number, slide in enumerate(prs.slides, start=1):
    if slide_number == 1:
        continue
    for sh in slide.shapes:
        if hasattr(sh, "text_frame") and "DERMTRACK AGENT" in sh.text and sh.text.strip()[:1].isdigit():
            sh.text_frame.paragraphs[0].text = f"{slide_number:02d}  /  DERMTRACK AGENT"
            break

prs.core_properties.title = "DermTrack Agent — Complete Review 1 Presentation"
prs.core_properties.subject = "Problem, literature, gap, objectives, solution, novelty, architecture, database, work plan"
prs.save(OUT)
print(OUT)
