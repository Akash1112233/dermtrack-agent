from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

OUT = "DermTrack_Introduction_Slide.pptx"
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
slide = prs.slides.add_slide(prs.slide_layouts[6])

# Palette
navy = RGBColor(10, 25, 47)
navy2 = RGBColor(16, 42, 67)
teal = RGBColor(35, 194, 176)
coral = RGBColor(255, 126, 103)
white = RGBColor(246, 250, 252)
muted = RGBColor(181, 204, 213)
line = RGBColor(52, 91, 111)

slide.background.fill.solid()
slide.background.fill.fore_color.rgb = navy

# Decorative right-side clinical-tech motif
for x, y, d, color, transparency in [
    (10.55, 0.58, 1.22, teal, 12),
    (11.25, 1.17, 1.95, teal, 40),
    (10.20, 4.95, 1.00, coral, 14),
    (11.00, 4.20, 1.60, coral, 48),
]:
    shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    shp.fill.solid(); shp.fill.fore_color.rgb = color; shp.fill.transparency = transparency
    shp.line.fill.background()

# Small connected-node motif
nodes = [(11.15, 2.20), (12.10, 1.55), (12.45, 2.65), (11.60, 3.25), (12.55, 4.05)]
for a, b in zip(nodes, nodes[1:]):
    connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(a[0]), Inches(a[1]), Inches(b[0]), Inches(b[1]))
    connector.line.color.rgb = line; connector.line.width = Pt(1.4)
for x, y in nodes:
    node = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x-0.07), Inches(y-0.07), Inches(0.14), Inches(0.14))
    node.fill.solid(); node.fill.fore_color.rgb = teal; node.line.fill.background()

# Accent block
accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.72), Inches(0.74), Inches(0.12), Inches(1.15))
accent.fill.solid(); accent.fill.fore_color.rgb = coral; accent.line.fill.background()

# Eyebrow
box = slide.shapes.add_textbox(Inches(1.08), Inches(0.72), Inches(8.4), Inches(0.35))
tf = box.text_frame; tf.clear(); tf.margin = 0
p = tf.paragraphs[0]; p.text = "PROJECT REVIEW PRESENTATION"; p.font.name = "Aptos"; p.font.size = Pt(13); p.font.bold = True; p.font.color.rgb = teal; p.font.letter_spacing = 2 if hasattr(p.font, 'letter_spacing') else None

# Main title
box = slide.shapes.add_textbox(Inches(1.05), Inches(1.30), Inches(8.8), Inches(1.05))
tf = box.text_frame; tf.clear(); tf.margin = 0; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
p = tf.paragraphs[0]; p.text = "DermTrack Agent"; p.font.name = "Aptos Display"; p.font.size = Pt(42); p.font.bold = True; p.font.color.rgb = white

# Subtitle
box = slide.shapes.add_textbox(Inches(1.08), Inches(2.42), Inches(8.4), Inches(0.72))
tf = box.text_frame; tf.clear(); tf.margin = 0
p = tf.paragraphs[0]; p.text = "AI-Assisted Non-Diagnostic\nDermatology Support System"; p.font.name = "Aptos"; p.font.size = Pt(21); p.font.color.rgb = muted; p.font.bold = False

# Bottom information panel
panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.82), Inches(4.12), Inches(8.92), Inches(2.25))
panel.fill.solid(); panel.fill.fore_color.rgb = navy2; panel.line.color.rgb = line; panel.line.width = Pt(1)

# Panel labels and values
items = [
    ("TEAM MEMBERS", "Annam Akash  ·  23MIA1112\nB. Hema  ·  23MIA1101"),
    ("PROJECT GUIDE", "Dr. Abhishi Choudhury"),
    ("DEPARTMENT", "Integrated M.Tech CSE (BA)"),
]
ys = [4.45, 4.45, 5.30]
xs = [1.20, 5.05, 1.20]
widths = [3.15, 3.95, 7.75]
for (label, value), x, y, w in zip(items, xs, ys, widths):
    lb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(0.23))
    tf = lb.text_frame; tf.clear(); tf.margin = 0
    p = tf.paragraphs[0]; p.text = label; p.font.name = "Aptos"; p.font.size = Pt(10); p.font.bold = True; p.font.color.rgb = teal
    vb = slide.shapes.add_textbox(Inches(x), Inches(y+0.30), Inches(w), Inches(0.58))
    tf = vb.text_frame; tf.clear(); tf.margin = 0
    p = tf.paragraphs[0]; p.text = value; p.font.name = "Aptos"; p.font.size = Pt(15); p.font.color.rgb = white

# Footer
footer = slide.shapes.add_textbox(Inches(10.55), Inches(6.72), Inches(2.1), Inches(0.25))
tf = footer.text_frame; tf.clear(); tf.margin = 0
p = tf.paragraphs[0]; p.text = "2026  /  INTRODUCTION"; p.font.name = "Aptos"; p.font.size = Pt(9); p.font.bold = True; p.font.color.rgb = muted; p.alignment = PP_ALIGN.RIGHT

prs.core_properties.title = "DermTrack Agent — Introduction"
prs.core_properties.subject = "Project review introduction slide"
prs.core_properties.author = "Annam Akash and B. Hema"
prs.save(OUT)
print(OUT)
